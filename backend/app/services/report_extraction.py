"""Medical-report text extraction.

Supports PDF (text-based or scanned), DOCX, and image reports. Scanned/image
content is read via local OCR (pytesseract), so no external API keys are
required — consistent with this project's "runs with zero API keys" ethos.

NOTE: the OCR path needs the Tesseract OCR binary on PATH, and scanned-PDF
rendering needs Poppler on PATH (pdf2image shells out to `pdftoppm`). Neither
is pip-installable. Text-based PDF/DOCX extraction works without them.
"""
import io
from typing import Optional

# Minimum average characters/page of native PDF text before we treat a PDF as
# "scanned" and fall back to OCR.
_MIN_CHARS_PER_PAGE = 20

SUPPORTED_EXTENSIONS = {"pdf", "docx", "jpg", "jpeg", "png"}


class OCRUnavailableError(RuntimeError):
    """Raised when OCR is needed but the Tesseract/Poppler binaries aren't
    present (e.g. a standard Vercel Python Function, which has no OS package
    manager — see README "Report analysis setup"). Text-based PDF/DOCX
    extraction doesn't need these and is unaffected.
    """


def _extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(filename: str, content_type: Optional[str], data: bytes) -> str:
    ext = _extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext or content_type or 'unknown'}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    if ext == "pdf":
        return _extract_pdf(data)
    if ext == "docx":
        return _extract_docx(data)
    return _extract_image(data)  # jpg/jpeg/png


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages_text = [(page.extract_text() or "") for page in reader.pages]
    native_text = "\n".join(pages_text).strip()

    avg_chars = len(native_text) / max(len(pages_text), 1)
    if avg_chars >= _MIN_CHARS_PER_PAGE:
        return native_text

    # Likely a scanned PDF with little/no embedded text — fall back to OCR.
    ocr_text = _ocr_pdf(data)
    return ocr_text or native_text


def _ocr_pdf(data: bytes) -> str:
    from pdf2image import convert_from_bytes
    from pdf2image.exceptions import PDFInfoNotInstalledError
    import pytesseract

    from app.config import get_settings

    settings = get_settings()
    try:
        images = convert_from_bytes(data)
        texts = [
            pytesseract.image_to_string(img, lang=settings.ocr_language) for img in images
        ]
    except (PDFInfoNotInstalledError, pytesseract.TesseractNotFoundError) as exc:
        raise OCRUnavailableError(
            "This looks like a scanned PDF, which needs OCR to read — OCR isn't "
            "available in this deployment. Please upload a text-based PDF or DOCX "
            "instead, or ask the site operator to enable OCR support."
        ) from exc
    return "\n".join(texts).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_image(data: bytes) -> str:
    from PIL import Image
    import pytesseract

    from app.config import get_settings

    settings = get_settings()
    image = Image.open(io.BytesIO(data))
    try:
        return pytesseract.image_to_string(image, lang=settings.ocr_language).strip()
    except pytesseract.TesseractNotFoundError as exc:
        raise OCRUnavailableError(
            "Reading image reports needs OCR, which isn't available in this "
            "deployment. Please upload a text-based PDF or DOCX instead, or ask "
            "the site operator to enable OCR support."
        ) from exc
