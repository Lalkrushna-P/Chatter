"""Medical-report text extraction tests (DOCX + unsupported-type handling).

PDF/image OCR paths need Tesseract/Poppler system binaries and are exercised
manually (see README "Report analysis setup"), not in this offline suite.
"""
import io

import docx
import pytest

from app.services.report_extraction import extract_text


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_docx_paragraphs():
    data = _make_docx_bytes(["Patient: Jane Doe", "Hemoglobin: 10.5 g/dL"])
    text = extract_text("report.docx", None, data)
    assert "Patient: Jane Doe" in text
    assert "Hemoglobin: 10.5 g/dL" in text


def test_extract_docx_table():
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Creatinine"
    table.rows[0].cells[1].text = "2.1 mg/dL"
    buf = io.BytesIO()
    document.save(buf)

    text = extract_text("labs.docx", None, buf.getvalue())
    assert "Creatinine" in text
    assert "2.1 mg/dL" in text


def test_unsupported_extension_raises():
    with pytest.raises(ValueError):
        extract_text("report.txt", "text/plain", b"some text")
